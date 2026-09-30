from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
d="/usr/share/fonts/truetype/dejavu/"
pdfmetrics.registerFont(TTFont("F",d+"DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("FB",d+"DejaVuSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("FI",d+"DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("FBI",d+"DejaVuSans-Bold.ttf"))
pdfmetrics.registerFontFamily("F",normal="F",bold="FB",italic="FI",boldItalic="FBI")

N=ParagraphStyle("n",fontName="F",fontSize=9.3,leading=13)
Q=ParagraphStyle("q",parent=N,fontName="FB",fontSize=9.8,leading=13.5,spaceBefore=9,spaceAfter=3,textColor=colors.HexColor("#1a3a6b"))
H=ParagraphStyle("h",parent=N,fontName="FB",fontSize=13,leading=16,textColor=colors.white,backColor=colors.HexColor("#1a3a6b"),borderPadding=(5,6,5,6),spaceBefore=8,spaceAfter=8)
T=ParagraphStyle("t",parent=N,fontName="FB",fontSize=17,leading=22,alignment=1,spaceAfter=2)
S=ParagraphStyle("s",parent=N,alignment=1,textColor=colors.grey,spaceAfter=6)
P=ParagraphStyle("p",parent=N,leftIndent=14,firstLineIndent=-14,spaceAfter=1.5)
EX=ParagraphStyle("e",parent=N,fontName="FI",leftIndent=14,textColor=colors.HexColor("#155724"),spaceBefore=2,spaceAfter=2)
SUB=ParagraphStyle("sub",parent=N,fontName="FB",spaceBefore=3,spaceAfter=1.5)
C=ParagraphStyle("c",parent=N,fontSize=8.6,leading=11.5)
CB=ParagraphStyle("cb",parent=C,fontName="FB",textColor=colors.white)

st=[]
def q(t): st.append(Paragraph(t,Q))
def sub(t): st.append(Paragraph(t,SUB))
def pts(lst,start=1):
    for i,t in enumerate(lst,start): st.append(Paragraph(f"<b>{i}.</b> {t}",P))
def ex(t): st.append(Paragraph("Example: "+t,EX))
def table(rows,widths):
    data=[[Paragraph(c,CB if r==0 else C) for c in row] for r,row in enumerate(rows)]
    t=Table(data,colWidths=[w*mm for w in widths],repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1a3a6b")),("GRID",(0,0),(-1,-1),.5,colors.grey),
      ("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#eef2f9")]),
      ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3)]))
    t._raw=rows
    st.append(Spacer(1,3)); st.append(t); st.append(Spacer(1,3))
def num(rows):
    data=[[Paragraph(f"<b>{a}</b>",C),Paragraph(b,C)] for a,b in rows]
    t=Table(data,colWidths=[27*mm,143*mm])
    t.setStyle(TableStyle([("BOX",(0,0),(-1,-1),1,colors.HexColor("#1a3a6b")),("LINEBELOW",(0,0),(-1,-2),.3,colors.lightgrey),
      ("BACKGROUND",(0,0),(0,-1),colors.HexColor("#eef2f9")),("VALIGN",(0,0),(-1,-1),"TOP")]))
    t._raw=[[f'<b>{a}</b>',b] for a,b in rows]
    st.append(Spacer(1,3)); st.append(KeepTogether(t)); st.append(Spacer(1,3))

st+= [Paragraph("SMLDS – Internal Exam Study Notes",T),Paragraph("Question Bank 02 · Section A (2 marks, 4 points each) · Section B (5 marks, 10 points each)",S),Paragraph("SECTION A – 2 Marks Each",H)]

q("Q1. What is the difference between precision and recall? Mention one situation where recall is preferred over precision.")
table([["Point","Precision","Recall"],
["1. Definition","Precision is the fraction of the instances predicted as positive that are actually positive, i.e. TP / (TP + FP).","Recall is the fraction of all actual positive instances that the model correctly identifies, i.e. TP / (TP + FN)."],
["2. What it tells us","It tells us how reliable the model's positive predictions are.","It tells us how many of the real positive cases the model is able to find."],
["3. Error it penalises","It is reduced by false positives, so it penalises false alarms.","It is reduced by false negatives, so it penalises missed cases."],
["4. When it matters","It is important when a false alarm is costly, e.g. spam filtering.","It is important when missing a positive is costly, e.g. disease detection."]],[30,70,70])
sub("Recall is preferred when:")
st.append(Paragraph("Cancer screening – missing a sick patient (FN) is far worse than a false alarm.",P))
ex("out of 100 cancer patients the model detects 95, so recall = 0.95.")

q("Q2. A classification model produces TP = 40, FP = 10, FN = 5, TN = 45. Calculate the precision.")
num([("Given","TP = 40, FP = 10, FN = 5, TN = 45"),("Formula","Precision = TP / (TP + FP)"),
("Substitution","Precision = 40 / (40 + 10) = 40 / 50"),("Result","<b>Precision = 0.80 (80%)</b>"),
("Conclusion","80% of the instances predicted positive are actually positive."),
("Extra check","Recall = 40/45 = 0.889; Accuracy = 85/100 = 0.85; F1 ≈ 0.842")])

q("Q3. What is SMOTE? Why is it used when handling an imbalanced classification dataset?")
pts(["SMOTE = <b>Synthetic Minority Over-sampling Technique</b>.",
"For a minority point it picks one of its k nearest minority neighbours and creates a synthetic point on the line between them: x_new = x + λ(x_neighbour − x), λ ∈ [0, 1].",
"Used because imbalanced data biases the model towards the majority class – accuracy looks high but minority recall is poor.",
"Unlike random duplication it creates new points, balancing classes and reducing overfitting. Apply on the training set only."])
ex("950 genuine and 50 fraud transactions – SMOTE generates synthetic fraud samples until both classes have about 950.")

q("Q4. What is the kernel trick in SVM? State its purpose.")
pts(["Some data is not linearly separable in its original space.",
"The kernel trick computes the dot product φ(x)·φ(x′) in a higher-dimensional space directly as K(x, x′), without computing the mapping φ.",
"<b>Purpose:</b> lets SVM draw non-linear decision boundaries efficiently and cheaply.",
"Common kernels: linear, polynomial (xᵀx′ + c)^d, RBF exp(−γ‖x−x′‖²), sigmoid."])
ex("one class forms a ring around another in 2-D; the RBF kernel lifts the points to a higher dimension where a plane separates them.")

q("Q5. What is the purpose of explained variance in PCA? How does it help in selecting the number of principal components?")
pts(["Explained variance = proportion of total variance captured by each principal component: λᵢ / Σλ.",
"It shows how much information each component retains.",
"Plot cumulative explained variance (scree plot); pick the smallest number of components reaching a threshold (90–95%) or use the elbow.",
"This keeps most information while reducing dimensions – simpler, faster, less noisy model."])
ex("PC1 = 60%, PC2 = 25%, PC3 = 10%, PC4 = 5% → cumulative 60, 85, 95, 100. To keep 95% choose 3 components.")

q("Q6. Mean = 50, SD = 5. Calculate the Z-score of an observation of 65. Is it an outlier using |Z| > 3?")
num([("Given","μ = 50, σ = 5, x = 65"),("Formula","Z = (x − μ) / σ"),("Substitution","Z = (65 − 50) / 5 = 15 / 5"),
("Result","<b>Z = 3</b>"),("Conclusion","Criterion is |Z| > 3. Here |Z| = 3, which is not greater than 3, so the observation is <b>not an outlier</b> (exactly on the borderline).")])

q("Q7. What is the main difference between PCA and LDA in terms of whether they use class labels?")
table([["Point","PCA","LDA"],
["1. Use of class labels","PCA does not use class labels; it is an unsupervised technique.","LDA uses the class labels; it is a supervised technique."],
["2. Objective","PCA finds the directions along which the variance of the data is maximum.","LDA finds the directions that maximise the separation between the classes."],
["3. Method","PCA takes the eigenvectors of the covariance matrix of the data.","LDA maximises the ratio of between-class scatter to within-class scatter."],
["4. Number of components","PCA can give as many components as there are features.","LDA can give at most (number of classes − 1) components."]],[30,70,70])
ex("for 3 flower species LDA gives at most 2 components that separate the species; PCA gives directions of maximum spread regardless of species.")

q("Q8. What is the silhouette score? What does a high score indicate?")
pts(["For each point s = (b − a) / max(a, b); a = mean distance to own cluster, b = mean distance to the nearest other cluster.",
"Value ranges from −1 to +1.",
"A high score (close to +1) means tight clusters, well separated from others – good clustering quality.",
"Near 0 = overlapping clusters; negative = likely wrong assignment. Also used to choose the best k."])
ex("a = 2, b = 8 → s = (8 − 2)/8 = 0.75, a well-separated point.")

q("Q9. Differentiate between AIC and BIC. Mention one key difference in how they penalise complexity.")
table([["Point","AIC","BIC"],
["1. Full form","AIC stands for Akaike Information Criterion.","BIC stands for Bayesian Information Criterion."],
["2. Formula","AIC = 2k − 2ln(L).","BIC = k·ln(n) − 2ln(L)."],
["3. <b>Penalty on complexity</b>","AIC charges a fixed penalty of 2 for every parameter added, whatever the sample size.","BIC charges a penalty of ln(n) per parameter, which increases as the sample size n grows."],
["4. Model preference","AIC tends to select a slightly more complex model and is better for prediction.","BIC penalises complexity more heavily (for n &gt; 7), so it selects simpler models and is better for finding the true model."]],[35,67,68])
st.append(Paragraph("Lower value is better for both (k = parameters, n = sample size, L = likelihood).",N))
ex("n = 100 → ln(100) ≈ 4.6, so BIC charges ≈ 4.6 per parameter vs AIC's 2; BIC rejects extra variables more readily.")

q("Q10. High training accuracy but much lower validation accuracy – identify the issue using bias–variance trade-off.")
pts(["Likely issue: <b>overfitting</b> – low bias, <b>high variance</b>.",
"The model memorised noise and specifics of the training data, so it generalises poorly.",
"The gap between training and validation accuracy is the sign of high variance.",
"Fixes: regularisation (L1/L2), more data, simpler model, pruning, dropout, cross-validation, early stopping."])
ex("training accuracy 98%, validation accuracy 70% → overfitting.")

st.append(Paragraph("SECTION B – 5 Marks Each (10 points each)",H))

q("B1. Fraud detection on an imbalanced dataset. a) Why can accuracy mislead? b) How do precision, recall and F1 help? c) Which metric to prioritise for detecting as many frauds as possible?")
sub("a) Why accuracy misleads")
pts(["Fraud may be under 1% of all transactions.","A useless model that predicts “genuine” for everything still scores ≈ 99% accuracy.","Accuracy counts all correct predictions equally, so the majority class dominates and it says nothing about fraud detection."])
ex("990 genuine + 10 fraud: always predicting “genuine” gives 99% accuracy but 0% recall.")
sub("b) Precision, recall, F1")
pts(["<b>Precision</b> = TP/(TP+FP): how many flagged transactions are truly fraud; low precision = many genuine customers disturbed.",
"<b>Recall</b> = TP/(TP+FN): how many real frauds are caught; exposes the misses accuracy hides.",
"<b>F1</b> = 2PR/(P+R): harmonic mean, high only when both are high – a single balanced score.",
"They focus on the positive (fraud) class so they reflect real performance on imbalanced data; confusion matrix / PR-curve adds detail."],4)
sub("c) Metric to prioritise")
pts(["Prioritise <b>recall</b>.","A missed fraud (FN) causes direct financial loss – far costlier than a false alarm needing a manual check.","Recall alone can be gamed by flagging everything, so also monitor precision; tune threshold or use F2."],8)

q("B2. Decision Tree: 95% training, 78% testing accuracy. a) Problem? b) How does pruning help? c) One advantage of Random Forest over a single tree.")
sub("a) Problem")
pts(["<b>Overfitting</b> (high variance).","The tree grows deep and memorises training noise.","The 17-point gap between train and test accuracy confirms it."])
sub("b) Pruning")
pts(["Pruning removes branches with little predictive power, simplifying the tree.",
"<i>Pre-pruning</i> stops growth early: max_depth, min_samples_split, min_samples_leaf, minimum impurity decrease.",
"<i>Post-pruning</i> grows the full tree then cuts weak branches, e.g. cost-complexity pruning (ccp_alpha) chosen on validation data.",
"Smaller tree = lower variance, better generalisation, easier to interpret."],4)
ex("limiting max_depth from 20 to 5 may drop train accuracy to 88% but raise test accuracy to 85%.")
sub("c) Random Forest vs single tree")
table([["Point","Single Decision Tree","Random Forest"],
["8. Structure","A single decision tree is built on the whole training set.","A random forest combines many trees, each trained on a bootstrap sample with a random subset of features (bagging)."],
["9. Variance / overfitting","A single deep tree has high variance and overfits the training data easily.","A random forest has low variance, because averaging the trees' predictions cancels out their individual errors."],
["10. Robustness","A single tree is sensitive to noise and small changes in the data.","A random forest is robust to noise and outliers and also gives feature importance."]],[30,70,70])

q("B3. Working principle of SVM. a) Hyperplane b) Margin and why maximise it c) Purpose of kernel trick with an example kernel.")
sub("Working principle")
pts(["SVM is a supervised classifier that finds the decision boundary separating the classes with the widest possible gap."])
sub("a) Hyperplane")
pts(["A hyperplane is the decision boundary w·x + b = 0.","It is a line in 2-D, a plane in 3-D, a hyperplane in higher dimensions.","Points with w·x + b &gt; 0 go to one class, &lt; 0 to the other."],2)
ex("in 2-D, the line x₁ + x₂ − 5 = 0 separates the two classes.")
sub("b) Margin")
pts(["Margin = distance between the hyperplane and the closest points of each class; those points are the <b>support vectors</b>. Margin = 2/‖w‖.","Maximising the margin means minimising ‖w‖.","Larger margin → better generalisation, lower overfitting risk, robust to noise. Only support vectors define the boundary."],5)
sub("c) Kernel trick")
pts(["For non-linearly separable data the kernel implicitly maps it to a higher-dimensional space where it becomes separable, without computing the mapping.","Example: <b>RBF kernel</b> K(x, x′) = exp(−γ‖x − x′‖²) handles circular/complex boundaries; others: polynomial, linear.","Soft-margin SVM uses parameter C to trade off margin width against misclassification."],8)

q("B4. Correlated features. a) How PCA uses covariance matrix and eigen decomposition b) Explained variance c) Why keep the first few PCs.")
sub("a) Covariance matrix and eigen decomposition")
pts(["Standardise the data (zero mean, unit variance).","Compute covariance matrix C = XᵀX/(n−1); correlated features give large off-diagonal values.","Eigen decomposition C v = λ v: eigenvectors = principal component directions, eigenvalues = variance along them.","Sort eigenvectors by eigenvalue (descending) and select the top k.","Project: Z = X·W_k – new components are uncorrelated and fewer."])
ex("10 correlated features reduced to 3 components.")
sub("b) Explained variance")
pts(["Ratio λᵢ / Σλ – share of total variance carried by each PC.","Cumulative sum (e.g. ≥ 95%) tells how many components to keep."],6)
ex("eigenvalues 4, 1, 0.5, 0.5 → PC1 = 4/6 = 66.7%.")
sub("c) Why keep the first few PCs")
pts(["Keep most information and remove redundancy / multicollinearity.","Reduce overfitting and the curse of dimensionality; speed up training.","Filter noise and allow 2-D / 3-D visualisation."],8)

q("B5. Customer transactions. a) K-means for customer groups b) One limitation of K-means and how DBSCAN addresses it c) Silhouette Score for evaluation.")
sub("a) K-means for customer groups")
pts(["Choose k using the elbow method or silhouette score.","Initialise k centroids (k-means++).","Assign each customer to the nearest centroid (Euclidean distance).","Recompute centroids as the mean of members; repeat until assignments stop changing. Scale features first."])
ex("k = 3 on (annual spend, visit frequency) → high-value, regular and occasional customers.")
sub("b) Limitation and DBSCAN")
table([["Point","K-means","DBSCAN"],
["5. Number of clusters","K-means requires the number of clusters k to be specified in advance.","DBSCAN finds the number of clusters automatically from the density of the data."],
["6. Cluster shape","K-means assumes clusters are roughly spherical and of similar size.","DBSCAN can find clusters of arbitrary shape."],
["7. Outliers","K-means is sensitive to outliers, as every point is forced into a cluster and pulls the centroids.","DBSCAN labels low-density points as noise (using eps and minPts), so outliers do not distort the clusters."]],[30,70,70])
sub("c) Silhouette score")
pts(["s = (b − a) / max(a, b); a = mean intra-cluster distance, b = mean nearest-cluster distance.","Average over all points, −1 to +1: near +1 good, near 0 overlap, negative misassignment.","Compute for different k and pick the k with the highest score."],8)
ex("average silhouette 0.42 (k=2), 0.61 (k=3), 0.48 (k=4) → choose k = 3.")

q("B6. High-dimensional data with outliers. a) t-SNE / UMAP b) IQR method c) One tree-based outlier detection method.")
sub("a) t-SNE / UMAP")
pts(["Both are non-linear dimensionality-reduction methods projecting high-dimensional data to 2-D/3-D for visualisation."])
ex("784-pixel handwritten digits plotted in 2-D show ten clusters, one per digit.")
table([["Point","t-SNE","UMAP"],
["2. Idea","t-SNE preserves the local neighbourhoods of the data, using probability distributions and a t-distribution in the low-dimensional space.","UMAP is based on manifold learning and builds a neighbour graph, which it then lays out in low dimensions."],
["3. Speed and structure","t-SNE is slower on large data and preserves global structure less well.","UMAP is faster and preserves more of the global structure along with the local structure."]],[30,70,70])
sub("b) IQR method")
pts(["Compute Q1 (25th percentile) and Q3 (75th percentile).","IQR = Q3 − Q1.","Lower fence = Q1 − 1.5×IQR; Upper fence = Q3 + 1.5×IQR.","Values outside the fences are outliers (same rule as boxplot whiskers); no normality assumption; applied per feature."],4)
ex("Q1 = 20, Q3 = 40 → IQR = 20, fences = −10 and 70; a value of 95 is an outlier.")
sub("c) Isolation Forest")
pts(["Outliers are few and different, so they are easy to <i>isolate</i>.","Builds random trees using random features and random split values; outliers are isolated in <b>fewer splits</b> (short path length).","Anomaly score is based on the path length averaged over all trees – a shorter average path means the point is more likely an anomaly."],8)

st.append(Paragraph("Sources cross-checked",H))
for t in ["scikit-learn docs – silhouette_score / silhouette_samples: s = (b − a)/max(a, b), range −1 to 1 (scikit-learn.org)",
"scikit-learn docs – IsolationForest: random feature + random split; outliers need fewer splits (shorter path length) (scikit-learn.org)",
"scikit-learn docs – PCA: explained_variance_ratio_ = eigenvalue / sum of eigenvalues; n_components=0.95 keeps 95% variance (scikit-learn.org)",
"scikit-learn docs – DBSCAN: core samples, eps, min_samples, noise label −1; arbitrary shapes, no k needed (scikit-learn.org)",
"SMOTE (Chawla et al., 2002, arxiv.org/abs/1106.1813): x_new = x + δ(x_neighbour − x), δ random in [0,1]; also GeeksforGeeks “SMOTE for Imbalanced Classification”",
"AIC vs BIC: penalty 2k vs k·ln(n); BIC stricter for n &gt; 7 (CASRAI comparison, standard textbooks)",
"LDA vs PCA: LDA supervised, at most (classes − 1) components; PCA unsupervised (scikit-learn LDA/PCA docs and tutorials)"]:
    st.append(Paragraph("• "+t,P))
st.append(Paragraph("Note: verified via search-result summaries of these pages; the pages themselves could not be opened in this environment. Confirm wording against your class notes.",EX))

st_copy=list(st)
def foot(c,doc):
    c.setFont("F",8);c.setFillColor(colors.grey);c.drawCentredString(A4[0]/2,10*mm,f"SMLDS Study Notes – page {doc.page}")
SimpleDocTemplate("/home/user/cohort/smlds/SMLDS_Study_Notes.pdf",pagesize=A4,leftMargin=20*mm,rightMargin=20*mm,topMargin=15*mm,bottomMargin=18*mm,title="SMLDS Study Notes").build(st,onFirstPage=foot,onLaterPages=foot)

# ---- editable Word version built from the same content ----
import re, html
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.platypus import KeepTogether as _KT

doc=Document()
for sec in doc.sections: sec.left_margin=sec.right_margin=Cm(2); sec.top_margin=sec.bottom_margin=Cm(1.8)
doc.styles["Normal"].font.name="Calibri"; doc.styles["Normal"].font.size=Pt(11)
NAVY=RGBColor(0x1a,0x3a,0x6b)
def shade(cell,hexcol):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd")
    sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexcol); tcPr.append(sh)
def runs(par,markup,bold=False,italic=False,color=None,size=None):
    for tok in re.split(r"(</?[bi]>)",markup):
        if tok=="<b>": bold=True
        elif tok=="</b>": bold=False
        elif tok=="<i>": italic=True
        elif tok=="</i>": italic=False
        elif tok:
            r=par.add_run(html.unescape(tok)); r.bold=bold; r.italic=italic
            if color: r.font.color.rgb=color
            if size: r.font.size=Pt(size)
def add_par(p):
    name=p.style.name; t=p.text
    if name=="t":
        par=doc.add_paragraph(); par.alignment=1; runs(par,t,bold=True,size=20)
    elif name=="s":
        par=doc.add_paragraph(); par.alignment=1; runs(par,t,color=RGBColor(0x66,0x66,0x66),size=10)
    elif name=="h":
        par=doc.add_paragraph(); runs(par,t,bold=True,color=NAVY,size=14)
        pPr=par._p.get_or_add_pPr(); b=OxmlElement("w:pBdr"); bt=OxmlElement("w:bottom")
        for k,v in (("val","single"),("sz","12"),("space","1"),("color","1A3A6B")): bt.set(qn("w:"+k),v)
        b.append(bt); pPr.append(b); par.paragraph_format.space_before=Pt(12)
    elif name=="q":
        par=doc.add_paragraph(); runs(par,t,bold=True,color=NAVY,size=11.5); par.paragraph_format.space_before=Pt(10); par.paragraph_format.keep_with_next=True
    elif name=="sub":
        par=doc.add_paragraph(); runs(par,t,bold=True); par.paragraph_format.keep_with_next=True
    elif name=="e":
        par=doc.add_paragraph(); runs(par,t,italic=True,color=RGBColor(0x15,0x57,0x24)); par.paragraph_format.left_indent=Cm(0.6)
    elif name=="p":
        par=doc.add_paragraph(); runs(par,t); par.paragraph_format.left_indent=Cm(0.6); par.paragraph_format.first_line_indent=Cm(-0.6); par.paragraph_format.space_after=Pt(2)
    else:
        par=doc.add_paragraph(); runs(par,t)
def add_table(tb):
    rows=tb._raw; nc=len(rows[0]); is_num=nc==2
    t=doc.add_table(rows=len(rows),cols=nc); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,row in enumerate(rows):
        for j,cp in enumerate(row):
            c=t.cell(i,j); c.text=""; par=c.paragraphs[0]
            if is_num:
                runs(par,cp); 
                if j==0: shade(c,"EEF2F9")
            elif i==0:
                runs(par,cp,bold=True,color=RGBColor(255,255,255)); shade(c,"1A3A6B")
            else:
                runs(par,cp)
                if i%2==0: shade(c,"EEF2F9")
            for r in par.runs: r.font.size=Pt(10)
    doc.add_paragraph()
def walk(items):
    for x in items:
        if isinstance(x,Paragraph): add_par(x)
        elif isinstance(x,Table): add_table(x)
        elif isinstance(x,_KT): walk(x._content)
walk(st_copy)
doc.save("/home/user/cohort/smlds/SMLDS_Study_Notes.docx")
